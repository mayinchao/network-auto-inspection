import csv
import logging
import re
import socket
import time
from datetime import datetime
from getpass import getpass
from pathlib import Path

import paramiko
from netmiko import ConnectHandler
from netmiko.exceptions import (
    NetmikoAuthenticationException,
    NetmikoTimeoutException,
)
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


# ============================================================
# 项目目录
# ============================================================

DEVICE_FILE = Path("config/devices.csv")

OUTPUT_DIR = Path("output")
BACKUP_DIR = Path("backup")
LOG_DIR = Path("logs")
REPORT_DIR = Path("reports")
DEBUG_DIR = OUTPUT_DIR / "debug"


# ============================================================
# SSH / Netmiko 巡检命令
# ============================================================

COMMANDS = {
    "windows": [
        "hostname",
        "whoami",
        "ipconfig",
    ],

    "huawei_vrp": [
        "display version",
        "display ip interface brief",
        "display vlan",
        "display ip routing-table",
        "display ospf peer",
    ],
}


BACKUP_COMMANDS = {
    "windows": "ipconfig /all",
    "huawei_vrp": "display current-configuration",
}


# ============================================================
# eNSP Huawei Console 巡检命令
# ============================================================

HUAWEI_CONSOLE_COMMANDS = [
    {
        "command": "display version",
        "required_any": [
            "Huawei Versatile Routing Platform",
            "VRP (R) software",
            "VRP",
        ],
        "timeout": 30,
    },
    {
        "command": "display ip interface brief",
        "required_any": [
            "IP Address/Mask",
            "Physical",
            "Protocol",
        ],
        "timeout": 30,
    },
    {
        "command": "display ip routing-table",
        "required_any": [
            "Routing Tables",
            "Destination/Mask",
            "Route Flags",
        ],
        "timeout": 30,
    },
]


MORE_PATTERN = re.compile(
    r"-+\s*More\s*-+",
    re.IGNORECASE,
)

ANSI_PATTERN = re.compile(
    r"\x1b\[[0-9;?]*[A-Za-z]"
)


# ============================================================
# 日志
# ============================================================


def setup_logging():
    LOG_DIR.mkdir(exist_ok=True)

    logger = logging.getLogger(
        "network_inspection"
    )

    logger.setLevel(logging.INFO)

    if logger.handlers:
        return logger

    formatter = logging.Formatter(
        "%(asctime)s | %(levelname)s | %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
    )

    file_handler = logging.FileHandler(
        LOG_DIR / "inspection.log",
        encoding="utf-8",
    )
    file_handler.setFormatter(formatter)

    console_handler = logging.StreamHandler()
    console_handler.setFormatter(formatter)

    logger.addHandler(file_handler)
    logger.addHandler(console_handler)

    return logger


logger = setup_logging()


# ============================================================
# Windows 命令编码处理
# ============================================================


def decode_output(data: bytes) -> str:
    for encoding in (
        "utf-8",
        "gbk",
    ):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue

    return data.decode(
        "utf-8",
        errors="replace",
    )


# ============================================================
# 读取设备清单
# ============================================================


def load_devices():
    devices = []

    with DEVICE_FILE.open(
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as file:
        reader = csv.DictReader(file)

        for row in reader:
            row["port"] = int(row["port"])

            # 可选健康检查字段。
            # 当前旧版 devices.csv 没这些列也不会报错。
            row["expected_interface"] = (
                row.get("expected_interface") or ""
            ).strip()
            row["expected_ip"] = (
                row.get("expected_ip") or ""
            ).strip()
            row["expected_route"] = (
                row.get("expected_route") or ""
            ).strip()

            # 当前 AR1 的实验环境默认值。
            # 这样旧版 5 列 CSV 只要新增 AR1 行就可以运行。
            if (
                row["device_type"] == "huawei_console"
                and row["name"] == "AR1"
            ):
                row["expected_interface"] = (
                    row["expected_interface"]
                    or "Ethernet0/0/8"
                )
                row["expected_ip"] = (
                    row["expected_ip"]
                    or "192.168.100.2/24"
                )
                row["expected_route"] = (
                    row["expected_route"]
                    or "192.168.100.0/24"
                )

            devices.append(row)

    return devices


# ============================================================
# Windows SSH 连接
# ============================================================


def connect_windows(device, password):
    client = paramiko.SSHClient()

    client.set_missing_host_key_policy(
        paramiko.AutoAddPolicy()
    )

    client.connect(
        hostname=device["host"],
        port=device["port"],
        username=device["username"],
        password=password,
        timeout=10,
        auth_timeout=10,
        banner_timeout=10,
        look_for_keys=False,
        allow_agent=False,
    )

    return client


# ============================================================
# Huawei VRP SSH / Netmiko 连接
# ============================================================


def connect_huawei_vrp(device, password):
    connection = ConnectHandler(
        device_type="huawei_vrp",
        host=device["host"],
        port=device["port"],
        username=device["username"],
        password=password,
        conn_timeout=10,
        banner_timeout=10,
        auth_timeout=10,
    )

    return connection


# ============================================================
# SSH / Netmiko 统一连接
# ============================================================


def connect_device(device, password):
    device_type = device["device_type"]

    if device_type == "windows":
        return (
            "paramiko",
            connect_windows(device, password),
        )

    if device_type == "huawei_vrp":
        return (
            "netmiko",
            connect_huawei_vrp(device, password),
        )

    raise ValueError(
        f"不支持的 SSH 设备类型：{device_type}"
    )


# ============================================================
# SSH / Netmiko 命令执行
# ============================================================


def run_command(
    connection_type,
    client,
    command,
):
    if connection_type == "paramiko":
        stdin, stdout, stderr = (
            client.exec_command(command)
        )

        output = decode_output(
            stdout.read()
        ).strip()

        error = decode_output(
            stderr.read()
        ).strip()

        return output, error

    if connection_type == "netmiko":
        output = client.send_command(
            command,
            read_timeout=60,
        )

        return output.strip(), ""

    raise ValueError("未知连接类型")


# ============================================================
# SSH / Netmiko 巡检
# ============================================================


def run_ssh_inspection(
    connection_type,
    client,
    device,
    timestamp,
):
    commands = COMMANDS.get(
        device["device_type"]
    )

    if commands is None:
        raise ValueError(
            "没有配置对应巡检命令"
        )

    sections = []

    for command in commands:
        logger.info(
            "设备 %s 执行巡检命令：%s",
            device["name"],
            command,
        )

        output, error = run_command(
            connection_type,
            client,
            command,
        )

        section = [
            "=" * 70,
            f"DEVICE: {device['name']}",
            f"HOST: {device['host']}",
            f"TYPE: {device['device_type']}",
            f"COMMAND: {command}",
            "=" * 70,
            output,
        ]

        if error:
            section.extend([
                "",
                "[STDERR]",
                error,
            ])

            logger.warning(
                "设备 %s 命令 %s 存在错误输出",
                device["name"],
                command,
            )

        sections.append("\n".join(section))

    combined_text = "\n\n".join(sections)

    OUTPUT_DIR.mkdir(exist_ok=True)

    output_file = (
        OUTPUT_DIR
        / f"{device['name']}_{timestamp}.txt"
    )

    output_file.write_text(
        combined_text,
        encoding="utf-8",
    )

    logger.info(
        "设备 %s 巡检结果保存至 %s",
        device["name"],
        output_file,
    )

    return output_file, combined_text


# ============================================================
# SSH / Netmiko 配置备份
# ============================================================


def backup_ssh_configuration(
    connection_type,
    client,
    device,
    timestamp,
):
    command = BACKUP_COMMANDS.get(
        device["device_type"]
    )

    if command is None:
        logger.warning(
            "设备 %s 没有配置备份命令",
            device["name"],
        )
        return None

    logger.info(
        "设备 %s 执行配置备份命令：%s",
        device["name"],
        command,
    )

    output, error = run_command(
        connection_type,
        client,
        command,
    )

    if error:
        logger.warning(
            "设备 %s 配置备份存在错误输出：%s",
            device["name"],
            error,
        )

    device_backup_dir = (
        BACKUP_DIR / device["name"]
    )
    device_backup_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    backup_file = (
        device_backup_dir
        / f"{device['name']}_{timestamp}.txt"
    )

    backup_file.write_text(
        output,
        encoding="utf-8",
    )

    logger.info(
        "设备 %s 配置备份保存至 %s",
        device["name"],
        backup_file,
    )

    return backup_file


# ============================================================
# 关闭 SSH / Netmiko 连接
# ============================================================


def close_connection(
    connection_type,
    client,
):
    if client is None:
        return

    try:
        if connection_type == "paramiko":
            client.close()

        elif connection_type == "netmiko":
            client.disconnect()

    except Exception as error:
        logger.warning(
            "关闭连接时出现异常：%s",
            error,
        )


# ============================================================
# Huawei Console 文本处理
# ============================================================


def normalize_terminal_output(text):
    text = ANSI_PATTERN.sub("", text)
    text = text.replace("\x00", "")
    text = text.replace("\r", "")
    return text



def clean_terminal_output(text):
    text = normalize_terminal_output(text)
    text = MORE_PATTERN.sub("", text)
    text = text.replace("\x08", "")
    text = text.replace("\x07", "")
    return text.strip()



def is_device_prompt(line, device_name):
    line = line.strip()
    return line in {
        f"<{device_name}>",
        f"[{device_name}]",
    }



def get_last_nonempty_line(text):
    lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]

    if not lines:
        return ""

    return lines[-1]



def drain_socket(sock, max_wait=0.3):
    old_timeout = sock.gettimeout()
    sock.settimeout(0.1)

    deadline = time.time() + max_wait
    data = ""

    while time.time() < deadline:
        try:
            chunk = sock.recv(4096)
            if not chunk:
                break

            data += chunk.decode(
                "utf-8",
                errors="ignore",
            )

        except socket.timeout:
            break

    sock.settimeout(old_timeout)
    return data



def wake_console(sock, device_name):
    old_data = drain_socket(
        sock,
        max_wait=0.5,
    )

    if old_data.strip():
        logger.info(
            "设备 %s 已清理 Console 历史残留数据",
            device_name,
        )

    sock.sendall(b"\r\n")
    time.sleep(0.3)

    response = drain_socket(
        sock,
        max_wait=0.5,
    )

    cleaned = clean_terminal_output(response)
    last_line = get_last_nonempty_line(cleaned)

    if is_device_prompt(
        last_line,
        device_name,
    ):
        logger.info(
            "设备 %s 检测到 Huawei CLI：%s",
            device_name,
            last_line,
        )
    else:
        logger.info(
            "设备 %s Console未主动回显提示符，继续通过命令输出验证",
            device_name,
        )



def validate_output(
    text,
    required_any=None,
    required_all=None,
):
    if required_any:
        if not any(
            keyword in text
            for keyword in required_any
        ):
            return False

    if required_all:
        if not all(
            keyword in text
            for keyword in required_all
        ):
            return False

    return True



def receive_console_command_output(
    sock,
    device_name,
    command,
    required_any=None,
    required_all=None,
    timeout=30,
    quiet_period=1.2,
):
    raw_data = ""
    deadline = time.time() + timeout
    last_receive_time = None
    handled_more_count = 0

    sock.settimeout(0.3)

    while time.time() < deadline:
        try:
            chunk = sock.recv(4096)
            if not chunk:
                break

            last_receive_time = time.time()

            raw_data += chunk.decode(
                "utf-8",
                errors="ignore",
            )

            normalized = normalize_terminal_output(
                raw_data
            )

            current_more_count = len(
                MORE_PATTERN.findall(
                    normalized
                )
            )

            while (
                handled_more_count
                < current_more_count
            ):
                logger.info(
                    "设备 %s 检测到分页，自动继续",
                    device_name,
                )

                sock.sendall(b" ")
                handled_more_count += 1
                time.sleep(0.15)

            cleaned = clean_terminal_output(
                raw_data
            )

            valid = validate_output(
                cleaned,
                required_any=required_any,
                required_all=required_all,
            )

            if valid:
                last_line = get_last_nonempty_line(
                    cleaned
                )

                if is_device_prompt(
                    last_line,
                    device_name,
                ):
                    return cleaned

        except socket.timeout:
            if last_receive_time is None:
                continue

            cleaned = clean_terminal_output(
                raw_data
            )

            valid = validate_output(
                cleaned,
                required_any=required_any,
                required_all=required_all,
            )

            silent_time = (
                time.time()
                - last_receive_time
            )

            if (
                valid
                and silent_time >= quiet_period
            ):
                return cleaned

    cleaned = clean_terminal_output(raw_data)

    raise TimeoutError(
        f"命令执行失败或输出不完整：{command}\n"
        f"已收到内容：\n{cleaned[-1000:]}"
    )



def save_console_debug(
    device_name,
    command,
    content,
):
    DEBUG_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    safe_command = (
        command
        .replace(" ", "_")
        .replace("/", "_")
    )

    debug_file = (
        DEBUG_DIR
        / f"{device_name}_{safe_command}_{timestamp}.txt"
    )

    debug_file.write_text(
        content,
        encoding="utf-8",
    )

    return debug_file



def send_console_command(
    sock,
    device_name,
    command,
    required_any=None,
    required_all=None,
    timeout=30,
):
    drain_socket(
        sock,
        max_wait=0.2,
    )

    logger.info(
        "设备 %s 执行 Console 命令：%s",
        device_name,
        command,
    )

    sock.sendall(
        (command + "\r\n").encode(
            "utf-8"
        )
    )

    try:
        output = receive_console_command_output(
            sock,
            device_name,
            command,
            required_any=required_any,
            required_all=required_all,
            timeout=timeout,
        )

    except Exception as exc:
        debug_file = save_console_debug(
            device_name,
            command,
            str(exc),
        )

        raise RuntimeError(
            f"{exc}\n调试信息已保存：{debug_file}"
        ) from exc

    meaningful = output
    meaningful = meaningful.replace(
        command,
        "",
    )
    meaningful = meaningful.replace(
        f"<{device_name}>",
        "",
    )
    meaningful = meaningful.replace(
        f"[{device_name}]",
        "",
    )
    meaningful = meaningful.replace(
        "#",
        "",
    )
    meaningful = meaningful.strip()

    if len(meaningful) < 5:
        raise RuntimeError(
            f"命令 {command} 返回有效内容过短：{meaningful!r}"
        )

    return output


# ============================================================
# Huawei 状态解析
# ============================================================


def check_huawei_interface(
    text,
    expected_interface,
    expected_ip,
):
    if not expected_interface:
        return {
            "status": "UNKNOWN",
            "interface": "-",
            "ip": "-",
            "physical": "-",
            "protocol": "-",
            "detail": "未配置期望接口",
        }

    for raw_line in text.splitlines():
        line = raw_line.strip()

        if not line.startswith(
            expected_interface
        ):
            continue

        fields = line.split()

        if len(fields) < 4:
            return {
                "status": "WARNING",
                "interface": expected_interface,
                "ip": "UNKNOWN",
                "physical": "UNKNOWN",
                "protocol": "UNKNOWN",
                "detail": "接口信息字段不完整",
            }

        interface = fields[0]
        ip_address = fields[1]
        physical = fields[2]
        protocol = fields[3]

        problems = []

        if (
            expected_ip
            and ip_address != expected_ip
        ):
            problems.append(
                f"IP异常：当前={ip_address}，期望={expected_ip}"
            )

        if physical.lower() != "up":
            problems.append(
                f"Physical={physical}"
            )

        if not protocol.lower().startswith(
            "up"
        ):
            problems.append(
                f"Protocol={protocol}"
            )

        return {
            "status": (
                "WARNING"
                if problems
                else "NORMAL"
            ),
            "interface": interface,
            "ip": ip_address,
            "physical": physical,
            "protocol": protocol,
            "detail": (
                "；".join(problems)
                if problems
                else "接口IP正确，Physical/Protocol均为up"
            ),
        }

    return {
        "status": "WARNING",
        "interface": expected_interface,
        "ip": "UNKNOWN",
        "physical": "UNKNOWN",
        "protocol": "UNKNOWN",
        "detail": "未找到目标接口",
    }



def check_huawei_route(
    text,
    expected_route,
):
    if not expected_route:
        return {
            "status": "UNKNOWN",
            "route": "-",
            "route_type": "-",
            "detail": "未配置期望路由",
        }

    for raw_line in text.splitlines():
        line = raw_line.strip()

        if expected_route not in line:
            continue

        fields = line.split()
        route_type = "UNKNOWN"

        for field in fields:
            if field.lower() in {
                "direct",
                "static",
                "ospf",
                "rip",
                "isis",
                "bgp",
            }:
                route_type = field
                break

        return {
            "status": "NORMAL",
            "route": expected_route,
            "route_type": route_type,
            "detail": (
                f"目标路由存在，路由类型={route_type}"
            ),
        }

    return {
        "status": "WARNING",
        "route": expected_route,
        "route_type": "UNKNOWN",
        "detail": "目标路由不存在",
    }



def calculate_huawei_health(
    interface_result,
    route_result,
    expected_ip,
):
    # 如果没有配置健康检查目标，只标记成功采集，不评分。
    if (
        interface_result["status"] == "UNKNOWN"
        and route_result["status"] == "UNKNOWN"
    ):
        return None, "SUCCESS"

    score = 0

    if (
        expected_ip
        and interface_result["ip"] == expected_ip
    ):
        score += 20

    if (
        interface_result["physical"].lower()
        == "up"
    ):
        score += 30

    if (
        interface_result["protocol"].lower()
        .startswith("up")
    ):
        score += 20

    if route_result["status"] == "NORMAL":
        score += 30

    if score >= 90:
        status = "NORMAL"
    elif score >= 60:
        status = "WARNING"
    else:
        status = "CRITICAL"

    return score, status



def generate_huawei_suggestion(
    interface_result,
    route_result,
    expected_ip,
):
    suggestions = []

    if (
        expected_ip
        and interface_result["ip"] != expected_ip
    ):
        suggestions.append(
            "核对接口IP地址配置"
        )

    if (
        interface_result["physical"].lower()
        != "up"
    ):
        suggestions.append(
            "检查接口物理连接及shutdown状态"
        )

    if not (
        interface_result["protocol"].lower()
        .startswith("up")
    ):
        suggestions.append(
            "检查接口协议状态及相关配置"
        )

    if route_result["status"] == "WARNING":
        suggestions.append(
            "检查接口状态及路由表，确认目标网段是否存在"
        )

    if not suggestions:
        return "当前未发现异常"

    return "；".join(suggestions)


# ============================================================
# Huawei Console 配置完整性
# ============================================================


def check_console_configuration(
    config_output,
    device_name,
):
    problems = []

    if len(config_output) < 500:
        problems.append(
            "配置输出长度过短"
        )

    if (
        f"sysname {device_name}"
        not in config_output
    ):
        problems.append(
            f"未检测到 sysname {device_name}"
        )

    if "return" not in config_output:
        problems.append(
            "未检测到配置结束标志 return"
        )

    if MORE_PATTERN.search(
        config_output
    ):
        problems.append(
            "配置中仍存在分页提示"
        )

    return problems


# ============================================================
# Huawei Console 巡检
# ============================================================


def inspect_huawei_console(
    device,
    timestamp,
):
    logger.info(
        "开始连接设备 %s (%s:%s) | 类型=huawei_console",
        device["name"],
        device["host"],
        device["port"],
    )

    with socket.create_connection(
        (
            device["host"],
            device["port"],
        ),
        timeout=5,
    ) as sock:
        logger.info(
            "设备 %s Console TCP连接成功",
            device["name"],
        )

        wake_console(
            sock,
            device["name"],
        )

        sections = []

        for item in HUAWEI_CONSOLE_COMMANDS:
            command = item["command"]

            output = send_console_command(
                sock,
                device["name"],
                command,
                required_any=item.get(
                    "required_any"
                ),
                required_all=item.get(
                    "required_all"
                ),
                timeout=item.get(
                    "timeout",
                    30,
                ),
            )

            section = [
                "=" * 70,
                f"DEVICE: {device['name']}",
                f"CONSOLE: {device['host']}:{device['port']}",
                "TYPE: huawei_console",
                f"COMMAND: {command}",
                "=" * 70,
                output,
            ]

            sections.append(
                "\n".join(section)
            )

        combined_text = "\n\n".join(
            sections
        )

        OUTPUT_DIR.mkdir(
            parents=True,
            exist_ok=True,
        )

        output_file = (
            OUTPUT_DIR
            / f"{device['name']}_console_{timestamp}.txt"
        )

        output_file.write_text(
            combined_text,
            encoding="utf-8",
        )

        logger.info(
            "设备 %s 巡检结果保存至 %s",
            device["name"],
            output_file,
        )

        # ----------------------------------------------------
        # 配置备份
        # ----------------------------------------------------

        backup_command = (
            "display current-configuration"
        )

        config_output = send_console_command(
            sock,
            device["name"],
            backup_command,
            required_all=[
                f"sysname {device['name']}",
                "return",
            ],
            timeout=120,
        )

        config_problems = (
            check_console_configuration(
                config_output,
                device["name"],
            )
        )

        backup_ok = not config_problems
        backup_file = None

        if backup_ok:
            device_backup_dir = (
                BACKUP_DIR
                / device["name"]
            )
            device_backup_dir.mkdir(
                parents=True,
                exist_ok=True,
            )

            backup_file = (
                device_backup_dir
                / f"{device['name']}_{timestamp}.txt"
            )

            backup_file.write_text(
                config_output,
                encoding="utf-8",
            )

            logger.info(
                "设备 %s 配置备份保存至 %s",
                device["name"],
                backup_file,
            )

        else:
            logger.warning(
                "设备 %s 配置完整性检查未通过：%s",
                device["name"],
                "；".join(config_problems),
            )

    # --------------------------------------------------------
    # Console断开以后做状态解析
    # --------------------------------------------------------

    interface_result = check_huawei_interface(
        combined_text,
        device["expected_interface"],
        device["expected_ip"],
    )

    route_result = check_huawei_route(
        combined_text,
        device["expected_route"],
    )

    health_score, health_status = (
        calculate_huawei_health(
            interface_result,
            route_result,
            device["expected_ip"],
        )
    )

    suggestion = generate_huawei_suggestion(
        interface_result,
        route_result,
        device["expected_ip"],
    )

    if not backup_ok:
        if health_status in {
            "SUCCESS",
            "NORMAL",
        }:
            health_status = "WARNING"

        suggestion += (
            "；配置备份完整性检查未通过"
        )

    score_text = (
        f"{health_score}/100"
        if health_score is not None
        else "-"
    )

    message = (
        f"接口={interface_result['status']}；"
        f"路由={route_result['status']}；"
        f"健康评分={score_text}；"
        f"配置备份={'PASS' if backup_ok else 'WARNING'}；"
        f"建议={suggestion}"
    )

    return {
        "success": True,
        "status": health_status,
        "message": message,
        "interface": interface_result["interface"],
        "interface_status": interface_result["status"],
        "route_status": route_result["status"],
        "health_score": score_text,
    }


# ============================================================
# 单台 Windows / Huawei VRP SSH 巡检
# ============================================================


def inspect_ssh_device(
    device,
    password,
    timestamp,
):
    connection_type = None
    client = None

    try:
        logger.info(
            "开始连接设备 %s (%s:%s) | 类型=%s",
            device["name"],
            device["host"],
            device["port"],
            device["device_type"],
        )

        (
            connection_type,
            client,
        ) = connect_device(
            device,
            password,
        )

        logger.info(
            "设备 %s SSH连接成功",
            device["name"],
        )

        _, combined_text = run_ssh_inspection(
            connection_type,
            client,
            device,
            timestamp,
        )

        backup_ssh_configuration(
            connection_type,
            client,
            device,
            timestamp,
        )

        # Huawei真实SSH将来也可以复用同一健康解析器。
        if device["device_type"] == "huawei_vrp":
            interface_result = check_huawei_interface(
                combined_text,
                device["expected_interface"],
                device["expected_ip"],
            )

            route_result = check_huawei_route(
                combined_text,
                device["expected_route"],
            )

            health_score, health_status = (
                calculate_huawei_health(
                    interface_result,
                    route_result,
                    device["expected_ip"],
                )
            )

            score_text = (
                f"{health_score}/100"
                if health_score is not None
                else "-"
            )

            message = (
                f"接口={interface_result['status']}；"
                f"路由={route_result['status']}；"
                f"健康评分={score_text}"
            )

            return {
                "success": True,
                "status": health_status,
                "message": message,
                "interface": interface_result["interface"],
                "interface_status": interface_result["status"],
                "route_status": route_result["status"],
                "health_score": score_text,
            }

        return {
            "success": True,
            "status": "SUCCESS",
            "message": "巡检成功",
            "interface": "-",
            "interface_status": "-",
            "route_status": "-",
            "health_score": "-",
        }

    except paramiko.AuthenticationException:
        message = "Paramiko SSH认证失败"
        logger.error(
            "设备 %s %s",
            device["name"],
            message,
        )

    except NetmikoAuthenticationException:
        message = "Netmiko SSH认证失败"
        logger.error(
            "设备 %s %s",
            device["name"],
            message,
        )

    except NetmikoTimeoutException as error:
        message = (
            f"Netmiko连接超时：{error}"
        )
        logger.error(
            "设备 %s %s",
            device["name"],
            message,
        )

    except paramiko.SSHException as error:
        message = (
            f"SSH协议错误：{error}"
        )
        logger.error(
            "设备 %s %s",
            device["name"],
            message,
        )

    except Exception as error:
        message = (
            f"巡检失败：{error}"
        )
        logger.error(
            "设备 %s %s",
            device["name"],
            message,
        )

    finally:
        close_connection(
            connection_type,
            client,
        )

    return {
        "success": False,
        "status": "FAILED",
        "message": message,
        "interface": "-",
        "interface_status": "-",
        "route_status": "-",
        "health_score": "-",
    }


# ============================================================
# 统一单设备巡检入口
# ============================================================


def inspect_device(
    device,
    password,
    timestamp,
):
    try:
        if device["device_type"] == "huawei_console":
            return inspect_huawei_console(
                device,
                timestamp,
            )

        return inspect_ssh_device(
            device,
            password,
            timestamp,
        )

    except Exception as error:
        message = (
            f"巡检失败：{error}"
        )

        logger.error(
            "设备 %s %s",
            device["name"],
            message,
        )

        return {
            "success": False,
            "status": "FAILED",
            "message": message,
            "interface": "-",
            "interface_status": "-",
            "route_status": "-",
            "health_score": "-",
        }


# ============================================================
# CSV 报告
# ============================================================


def generate_csv_report(
    summary_rows,
    timestamp,
):
    REPORT_DIR.mkdir(exist_ok=True)

    report_file = (
        REPORT_DIR
        / f"inspection_summary_{timestamp}.csv"
    )

    fieldnames = [
        "time",
        "name",
        "host",
        "port",
        "device_type",
        "status",
        "interface",
        "interface_status",
        "route_status",
        "health_score",
        "message",
    ]

    with report_file.open(
        "w",
        encoding="utf-8-sig",
        newline="",
    ) as file:
        writer = csv.DictWriter(
            file,
            fieldnames=fieldnames,
        )

        writer.writeheader()
        writer.writerows(summary_rows)

    logger.info(
        "CSV巡检报告保存至 %s",
        report_file,
    )


# ============================================================
# Excel 样式
# ============================================================


def status_fill(status):
    if status in {"SUCCESS", "NORMAL"}:
        return "C6EFCE"

    if status == "WARNING":
        return "FFEB9C"

    if status == "CRITICAL":
        return "F4B183"

    if status == "FAILED":
        return "FFC7CE"

    return "FFFFFF"


# ============================================================
# Excel 报告
# ============================================================


def generate_excel_report(
    summary_rows,
    timestamp,
):
    REPORT_DIR.mkdir(exist_ok=True)

    report_file = (
        REPORT_DIR
        / f"inspection_summary_{timestamp}.xlsx"
    )

    workbook = Workbook()

    # --------------------------------------------------------
    # Sheet 1：巡检汇总
    # --------------------------------------------------------

    sheet = workbook.active
    sheet.title = "巡检汇总"

    sheet.merge_cells("A1:K1")
    sheet["A1"] = "网络设备自动巡检汇总报告"
    sheet["A1"].font = Font(
        bold=True,
        size=16,
    )
    sheet["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center",
    )
    sheet.row_dimensions[1].height = 28

    sheet["A2"] = "生成时间"
    sheet["B2"] = datetime.now().strftime(
        "%Y-%m-%d %H:%M:%S"
    )

    headers = [
        "巡检时间",
        "设备名称",
        "管理地址",
        "端口",
        "设备类型",
        "总体状态",
        "接口",
        "接口状态",
        "路由状态",
        "健康评分",
        "说明 / 建议",
    ]

    for column_index, header in enumerate(
        headers,
        start=1,
    ):
        cell = sheet.cell(
            row=4,
            column=column_index,
            value=header,
        )

        cell.font = Font(bold=True)
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="D9EAF7",
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    for row_index, row_data in enumerate(
        summary_rows,
        start=5,
    ):
        values = [
            row_data["time"],
            row_data["name"],
            row_data["host"],
            row_data["port"],
            row_data["device_type"],
            row_data["status"],
            row_data["interface"],
            row_data["interface_status"],
            row_data["route_status"],
            row_data["health_score"],
            row_data["message"],
        ]

        for column_index, value in enumerate(
            values,
            start=1,
        ):
            cell = sheet.cell(
                row=row_index,
                column=column_index,
                value=value,
            )

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

        status_cell = sheet.cell(
            row=row_index,
            column=6,
        )
        status_cell.font = Font(bold=True)
        status_cell.alignment = Alignment(
            horizontal="center",
        )
        status_cell.fill = PatternFill(
            fill_type="solid",
            fgColor=status_fill(
                row_data["status"]
            ),
        )

    sheet.freeze_panes = "A5"

    if summary_rows:
        sheet.auto_filter.ref = (
            f"A4:K{sheet.max_row}"
        )

    column_widths = {
        "A": 21,
        "B": 16,
        "C": 17,
        "D": 10,
        "E": 18,
        "F": 14,
        "G": 20,
        "H": 14,
        "I": 14,
        "J": 14,
        "K": 70,
    }

    for column, width in (
        column_widths.items()
    ):
        sheet.column_dimensions[
            column
        ].width = width

    # --------------------------------------------------------
    # Sheet 2：统计
    # --------------------------------------------------------

    stats_sheet = workbook.create_sheet(
        "统计"
    )

    stats_sheet.merge_cells("A1:B1")
    stats_sheet["A1"] = "巡检统计"
    stats_sheet["A1"].font = Font(
        bold=True,
        size=15,
    )
    stats_sheet["A1"].alignment = Alignment(
        horizontal="center",
    )

    total_count = len(summary_rows)

    normal_count = sum(
        1
        for row in summary_rows
        if row["status"]
        in {"SUCCESS", "NORMAL"}
    )

    warning_count = sum(
        1
        for row in summary_rows
        if row["status"] == "WARNING"
    )

    critical_count = sum(
        1
        for row in summary_rows
        if row["status"] == "CRITICAL"
    )

    failed_count = sum(
        1
        for row in summary_rows
        if row["status"] == "FAILED"
    )

    execution_success_count = (
        total_count - failed_count
    )

    execution_success_rate = (
        execution_success_count
        / total_count
        if total_count
        else 0
    )

    stats_data = [
        ("设备总数", total_count),
        ("正常 / 成功", normal_count),
        ("警告", warning_count),
        ("严重异常", critical_count),
        ("巡检失败", failed_count),
        (
            "巡检执行成功率",
            execution_success_rate,
        ),
    ]

    for row_index, (
        label,
        value,
    ) in enumerate(
        stats_data,
        start=3,
    ):
        stats_sheet.cell(
            row=row_index,
            column=1,
            value=label,
        )
        stats_sheet.cell(
            row=row_index,
            column=2,
            value=value,
        )
        stats_sheet.cell(
            row=row_index,
            column=1,
        ).font = Font(bold=True)

    stats_sheet["B8"].number_format = (
        "0.00%"
    )

    stats_sheet.column_dimensions[
        "A"
    ].width = 22
    stats_sheet.column_dimensions[
        "B"
    ].width = 20

    # --------------------------------------------------------
    # Sheet 3：异常设备
    # --------------------------------------------------------

    error_sheet = workbook.create_sheet(
        "异常设备"
    )

    error_headers = [
        "时间",
        "设备名称",
        "地址",
        "端口",
        "设备类型",
        "状态",
        "接口",
        "接口状态",
        "路由状态",
        "健康评分",
        "原因 / 建议",
    ]

    for column_index, header in enumerate(
        error_headers,
        start=1,
    ):
        cell = error_sheet.cell(
            row=1,
            column=column_index,
            value=header,
        )
        cell.font = Font(bold=True)
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="F4CCCC",
        )
        cell.alignment = Alignment(
            horizontal="center",
        )

    abnormal_rows = [
        row
        for row in summary_rows
        if row["status"]
        in {
            "WARNING",
            "CRITICAL",
            "FAILED",
        }
    ]

    for row_index, row_data in enumerate(
        abnormal_rows,
        start=2,
    ):
        values = [
            row_data["time"],
            row_data["name"],
            row_data["host"],
            row_data["port"],
            row_data["device_type"],
            row_data["status"],
            row_data["interface"],
            row_data["interface_status"],
            row_data["route_status"],
            row_data["health_score"],
            row_data["message"],
        ]

        for column_index, value in enumerate(
            values,
            start=1,
        ):
            cell = error_sheet.cell(
                row=row_index,
                column=column_index,
                value=value,
            )
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

        status_cell = error_sheet.cell(
            row=row_index,
            column=6,
        )
        status_cell.fill = PatternFill(
            fill_type="solid",
            fgColor=status_fill(
                row_data["status"]
            ),
        )

    error_sheet.freeze_panes = "A2"

    error_widths = {
        "A": 21,
        "B": 16,
        "C": 17,
        "D": 10,
        "E": 18,
        "F": 14,
        "G": 20,
        "H": 14,
        "I": 14,
        "J": 14,
        "K": 70,
    }

    for column, width in (
        error_widths.items()
    ):
        error_sheet.column_dimensions[
            column
        ].width = width

    workbook.save(report_file)

    logger.info(
        "Excel巡检报告保存至 %s",
        report_file,
    )


# ============================================================
# 主程序
# ============================================================


def main():
    devices = load_devices()

    if not devices:
        logger.warning("设备清单为空")
        return

    logger.info(
        "========== 开始网络设备巡检 =========="
    )

    logger.info(
        "本次读取到 %s 台设备",
        len(devices),
    )

    # 只有存在真正 SSH 设备时才询问密码。
    need_ssh_password = any(
        device["device_type"]
        in {"windows", "huawei_vrp"}
        for device in devices
    )

    password = (
        getpass("SSH password: ")
        if need_ssh_password
        else ""
    )

    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    summary_rows = []

    for device in devices:
        result = inspect_device(
            device,
            password,
            timestamp,
        )

        summary_rows.append({
            "time": datetime.now().strftime(
                "%Y-%m-%d %H:%M:%S"
            ),
            "name": device["name"],
            "host": device["host"],
            "port": device["port"],
            "device_type": (
                device["device_type"]
            ),
            "status": result["status"],
            "interface": result["interface"],
            "interface_status": (
                result["interface_status"]
            ),
            "route_status": (
                result["route_status"]
            ),
            "health_score": (
                result["health_score"]
            ),
            "message": result["message"],
        })

    generate_csv_report(
        summary_rows,
        timestamp,
    )

    generate_excel_report(
        summary_rows,
        timestamp,
    )

    normal_count = sum(
        1
        for row in summary_rows
        if row["status"]
        in {"SUCCESS", "NORMAL"}
    )

    warning_count = sum(
        1
        for row in summary_rows
        if row["status"] == "WARNING"
    )

    critical_count = sum(
        1
        for row in summary_rows
        if row["status"] == "CRITICAL"
    )

    failed_count = sum(
        1
        for row in summary_rows
        if row["status"] == "FAILED"
    )

    logger.info(
        "巡检结束 | 总数=%s | 正常/成功=%s | 警告=%s | 严重异常=%s | 失败=%s",
        len(summary_rows),
        normal_count,
        warning_count,
        critical_count,
        failed_count,
    )

    logger.info(
        "========== 本次巡检完成 =========="
    )


if __name__ == "__main__":
    main()