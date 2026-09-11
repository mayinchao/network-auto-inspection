from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

from huawei_parser import (
    get_latest_inspection_file,
    check_interface,
    check_route,
)


BASE_DIR = Path(__file__).resolve().parent.parent
REPORT_DIR = BASE_DIR / "reports"

DEVICE_NAME = "AR1"

EXPECTED_IP = "192.168.100.2/24"


# ============================================================
# 健康评分
# ============================================================

def calculate_health_score(
    interface_result,
    route_result,
):
    score = 0
    details = []

    # --------------------------------------------------------
    # IP地址：20分
    # --------------------------------------------------------

    if interface_result["ip"] == EXPECTED_IP:

        score += 20

        details.append(
            ("IP地址", "NORMAL", 20)
        )

    else:

        details.append(
            ("IP地址", "WARNING", 0)
        )

    # --------------------------------------------------------
    # Physical：30分
    # --------------------------------------------------------

    physical = (
        interface_result["physical"]
        .lower()
    )

    if physical == "up":

        score += 30

        details.append(
            ("Physical状态", "NORMAL", 30)
        )

    else:

        details.append(
            (
                "Physical状态",
                "WARNING",
                0,
            )
        )

    # --------------------------------------------------------
    # Protocol：20分
    # --------------------------------------------------------

    protocol = (
        interface_result["protocol"]
        .lower()
    )

    if protocol.startswith("up"):

        score += 20

        details.append(
            ("Protocol状态", "NORMAL", 20)
        )

    else:

        details.append(
            (
                "Protocol状态",
                "WARNING",
                0,
            )
        )

    # --------------------------------------------------------
    # 路由：30分
    # --------------------------------------------------------

    if route_result["status"] == "NORMAL":

        score += 30

        details.append(
            ("目标路由", "NORMAL", 30)
        )

    else:

        details.append(
            ("目标路由", "WARNING", 0)
        )

    return score, details


# ============================================================
# 根据分数判断总体状态
# ============================================================

def get_health_status(score):

    if score >= 90:
        return "NORMAL"

    if score >= 60:
        return "WARNING"

    return "CRITICAL"


# ============================================================
# 生成运维建议
# ============================================================

def generate_suggestions(
    interface_result,
    route_result,
):

    suggestions = []

    if (
        interface_result["physical"]
        .lower()
        != "up"
    ):

        suggestions.append(
            "检查接口物理连接及 shutdown 状态"
        )

    if not (
        interface_result["protocol"]
        .lower()
        .startswith("up")
    ):

        suggestions.append(
            "检查接口协议状态及相关配置"
        )

    if (
        interface_result["ip"]
        != EXPECTED_IP
    ):

        suggestions.append(
            "核对接口 IP 地址配置"
        )

    if route_result["status"] != "NORMAL":

        suggestions.append(
            "检查接口状态及路由表，确认目标网段是否存在"
        )

    if not suggestions:

        suggestions.append(
            "当前未发现异常"
        )

    return "；".join(
        suggestions
    )


# ============================================================
# Excel格式
# ============================================================

def set_column_widths(sheet):

    widths = {
        "A": 18,
        "B": 24,
        "C": 24,
        "D": 18,
        "E": 18,
        "F": 42,
    }

    for column, width in widths.items():

        sheet.column_dimensions[
            column
        ].width = width


def format_header(sheet):

    for cell in sheet[1]:

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )


def apply_status_style(cell):

    if cell.value == "NORMAL":

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="C6EFCE",
        )

    elif cell.value == "WARNING":

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="FFEB9C",
        )

    elif cell.value == "CRITICAL":

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="FFC7CE",
        )


# ============================================================
# 生成Excel
# ============================================================

def create_excel_report(
    source_file,
    interface_result,
    route_result,
    score,
    health_status,
    details,
    suggestions,
):

    REPORT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    timestamp = (
        datetime.now().strftime(
            "%Y%m%d_%H%M%S"
        )
    )

    output_file = (
        REPORT_DIR
        / (
            f"{DEVICE_NAME}_"
            f"health_report_"
            f"{timestamp}.xlsx"
        )
    )

    wb = Workbook()

    # ========================================================
    # Sheet 1：巡检汇总
    # ========================================================

    ws = wb.active

    ws.title = "巡检汇总"

    ws.append(
        [
            "设备名称",
            "巡检时间",
            "接口",
            "总体状态",
            "健康评分",
            "处理建议",
        ]
    )

    ws.append(
        [
            DEVICE_NAME,
            datetime.now().strftime(
                "%Y-%m-%d %H:%M:%S"
            ),
            interface_result[
                "interface"
            ],
            health_status,
            f"{score}/100",
            suggestions,
        ]
    )

    format_header(
        ws
    )

    set_column_widths(
        ws
    )

    apply_status_style(
        ws["D2"]
    )

    # ========================================================
    # Sheet 2：状态详情
    # ========================================================

    detail_ws = wb.create_sheet(
        "状态详情"
    )

    detail_ws.append(
        [
            "检查项",
            "当前值",
            "状态",
            "得分",
            "说明",
        ]
    )

    detail_ws.append(
        [
            "接口IP",
            interface_result["ip"],
            (
                "NORMAL"
                if interface_result["ip"]
                == EXPECTED_IP
                else "WARNING"
            ),
            (
                20
                if interface_result["ip"]
                == EXPECTED_IP
                else 0
            ),
            (
                f"期望IP："
                f"{EXPECTED_IP}"
            ),
        ]
    )

    detail_ws.append(
        [
            "Physical",
            interface_result[
                "physical"
            ],
            (
                "NORMAL"
                if interface_result[
                    "physical"
                ].lower()
                == "up"
                else "WARNING"
            ),
            (
                30
                if interface_result[
                    "physical"
                ].lower()
                == "up"
                else 0
            ),
            "接口物理层状态",
        ]
    )

    detail_ws.append(
        [
            "Protocol",
            interface_result[
                "protocol"
            ],
            (
                "NORMAL"
                if interface_result[
                    "protocol"
                ].lower()
                .startswith("up")
                else "WARNING"
            ),
            (
                20
                if interface_result[
                    "protocol"
                ].lower()
                .startswith("up")
                else 0
            ),
            "接口协议层状态",
        ]
    )

    detail_ws.append(
        [
            "目标路由",
            route_result[
                "route"
            ],
            route_result[
                "status"
            ],
            (
                30
                if route_result[
                    "status"
                ]
                == "NORMAL"
                else 0
            ),
            route_result[
                "detail"
            ],
        ]
    )

    format_header(
        detail_ws
    )

    for row in range(
        2,
        detail_ws.max_row + 1,
    ):

        apply_status_style(
            detail_ws[
                f"C{row}"
            ]
        )

    for column in range(
        1,
        6,
    ):

        detail_ws.column_dimensions[
            get_column_letter(
                column
            )
        ].width = 24

    detail_ws.column_dimensions[
        "E"
    ].width = 45

    # ========================================================
    # Sheet 3：报告信息
    # ========================================================

    info_ws = wb.create_sheet(
        "报告信息"
    )

    info_ws.append(
        [
            "项目",
            "内容",
        ]
    )

    info_ws.append(
        [
            "设备",
            DEVICE_NAME,
        ]
    )

    info_ws.append(
        [
            "原始巡检文件",
            source_file.name,
        ]
    )

    info_ws.append(
        [
            "健康评分",
            f"{score}/100",
        ]
    )

    info_ws.append(
        [
            "总体状态",
            health_status,
        ]
    )

    info_ws.append(
        [
            "处理建议",
            suggestions,
        ]
    )

    format_header(
        info_ws
    )

    info_ws.column_dimensions[
        "A"
    ].width = 22

    info_ws.column_dimensions[
        "B"
    ].width = 65

    wb.save(
        output_file
    )

    return output_file


# ============================================================
# 主程序
# ============================================================

def main():

    print(
        "=" * 60
    )

    print(
        "Huawei 网络健康报告生成"
    )

    print(
        "=" * 60
    )

    # --------------------------------------------------------
    # 获取最近一次巡检数据
    # --------------------------------------------------------

    latest_file = (
        get_latest_inspection_file()
    )

    print(
        f"数据来源："
        f"{latest_file.name}"
    )

    text = latest_file.read_text(
        encoding="utf-8",
        errors="ignore",
    )

    # --------------------------------------------------------
    # 状态解析
    # --------------------------------------------------------

    interface_result = (
        check_interface(
            text
        )
    )

    route_result = (
        check_route(
            text
        )
    )

    # --------------------------------------------------------
    # 健康评分
    # --------------------------------------------------------

    score, details = (
        calculate_health_score(
            interface_result,
            route_result,
        )
    )

    health_status = (
        get_health_status(
            score
        )
    )

    suggestions = (
        generate_suggestions(
            interface_result,
            route_result,
        )
    )

    # --------------------------------------------------------
    # 输出到屏幕
    # --------------------------------------------------------

    print()

    print(
        f"设备：{DEVICE_NAME}"
    )

    print(
        f"接口："
        f"{interface_result['interface']}"
    )

    print(
        f"IP："
        f"{interface_result['ip']}"
    )

    print(
        f"Physical："
        f"{interface_result['physical']}"
    )

    print(
        f"Protocol："
        f"{interface_result['protocol']}"
    )

    print(
        f"路由："
        f"{route_result['status']}"
    )

    print()

    print(
        f"健康评分："
        f"{score}/100"
    )

    print(
        f"总体状态："
        f"{health_status}"
    )

    print(
        f"处理建议："
        f"{suggestions}"
    )

    # --------------------------------------------------------
    # Excel
    # --------------------------------------------------------

    report_file = (
        create_excel_report(
            latest_file,
            interface_result,
            route_result,
            score,
            health_status,
            details,
            suggestions,
        )
    )

    print()

    print(
        f"Excel巡检报告已生成："
        f"{report_file}"
    )

    print(
        "=" * 60
    )


if __name__ == "__main__":

    main()